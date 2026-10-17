from __future__ import annotations

from io import BytesIO
import json
from pathlib import Path

from fastapi.testclient import TestClient
import pytest

from src.webapp.api import create_app
from src.webapp.export_service import render_export
from src.webapp.pdf_ingestion import ParsedPage


def _application(tmp_path):
    app = create_app(run_root=tmp_path / "science-runs", serve_frontend=False)
    app.state.job_runner.submit = lambda _job_id: None
    return app, TestClient(app)


def _polish_provider(monkeypatch, transform):
    calls = []

    def build(**settings):
        def invoke(messages, **kwargs):
            assert messages[0]["role"] == "system"
            assert messages[1]["role"] == "user"
            context = json.loads(messages[1]["content"])["task_context"]
            calls.append({"settings": settings, "context": context})
            segment = context["segments"][0]
            return {
                "schema_version": "polish_revision_v1",
                "segments": [{
                    "segment_id": segment["segment_id"], "original": segment["text"],
                    "revised": transform(segment["text"]), "rationale": "Clarify wording.",
                    "risk": "low", "preserved_entities": [], "unresolved_items": [],
                }],
                "global_notes": [],
            }
        return invoke

    monkeypatch.setattr("src.webapp.polish_service.build_web_json_llm_call", build)
    return calls


def test_polish_reviews_preserve_source_offsets_and_publish_exports(tmp_path: Path, monkeypatch):
    calls = _polish_provider(monkeypatch, lambda text: text.replace("was observed", "was measured"))
    app, client = _application(tmp_path)
    original = "  The response was observed at 5 mg.\n\n\nThe signal was observed at 8 mg.  \n"
    created = client.post("/api/polish/jobs", json={"text": original, "model_name": "qwen3.7-plus"})
    assert created.status_code == 202, created.text
    job_id = created.json()["job_id"]
    app.state.job_runner._execute(job_id)
    assert app.state.job_store.get(job_id).status == "WAITING_FOR_REVIEW", app.state.job_store.get(job_id).error_message
    assert len(calls) == 2
    assert {call["settings"]["model"] for call in calls} == {"qwen3.7-plus"}
    assert client.post(f"/api/polish/jobs/{job_id}/export", json={}).status_code == 409
    rejected = client.post(f"/api/polish/jobs/{job_id}/actions", json={
        "type": "reject_segments", "segment_ids": ["p-0002"], "expected_review_version": 0,
    })
    assert rejected.status_code == 202, rejected.text
    stale = client.post(f"/api/polish/jobs/{job_id}/actions", json={
        "type": "accept", "expected_review_version": 0,
    })
    assert stale.status_code == 409
    unknown = client.post(f"/api/polish/jobs/{job_id}/actions", json={
        "type": "accept_segments", "segment_ids": ["p-9999"],
    })
    assert unknown.status_code == 409
    accepted = client.post(f"/api/jobs/{job_id}/review", json={"accepted": True, "expected_review_version": 1})
    assert accepted.status_code == 202, accepted.text
    exported = client.post(f"/api/polish/jobs/{job_id}/export", json={"expected_review_version": 2})
    assert exported.status_code == 200, exported.text
    assert exported.text == original.replace("was observed", "was measured", 1)
    artifact_id = exported.headers["x-artifact-id"]
    assert artifact_id in app.state.job_store.get(job_id).output_artifact_ids
    assert client.get(f"/api/jobs/{job_id}/artifacts/{artifact_id}").content == exported.content
    repeated = client.post(f"/api/polish/jobs/{job_id}/export", json={})
    assert repeated.headers["x-artifact-id"] == artifact_id
    structured = client.post(f"/api/polish/jobs/{job_id}/export", json={"format": "json"}).json()
    assert structured["segments"][1]["review_status"] == "rejected"
    app.state.job_runner.shutdown()


def test_invalid_polish_is_not_saved_as_checkpoint(tmp_path: Path, monkeypatch):
    _polish_provider(monkeypatch, lambda text: text.replace("5 mg", "6 mg"))
    app, client = _application(tmp_path)
    created = client.post("/api/polish/jobs", json={"text": "The response was observed at 5 mg."})
    job_id = created.json()["job_id"]
    app.state.job_runner._execute(job_id)
    assert app.state.job_store.get(job_id).status == "FAILED"
    assert not list((app.state.job_store.job_directory(job_id) / "checkpoints").glob("*.json"))
    assert client.get(f"/api/jobs/{job_id}/artifacts").json() == []
    app.state.job_runner.shutdown()


def test_rejected_polish_does_not_expose_previous_attempt_result(tmp_path: Path, monkeypatch):
    _polish_provider(monkeypatch, lambda text: text.replace("observed", "measured"))
    app, client = _application(tmp_path)
    created = client.post("/api/polish/jobs", json={"text": "The response was observed."})
    job_id = created.json()["job_id"]
    app.state.job_runner._execute(job_id)
    rejected = client.post(
        f"/api/polish/jobs/{job_id}/actions",
        json={"type": "reject", "notes": "Revise the paragraph."},
    )
    assert rejected.status_code == 202
    assert app.state.job_store.get(job_id).status == "QUEUED"
    assert client.get(f"/api/polish/jobs/{job_id}/revisions").status_code == 409

    app.state.job_runner._execute(job_id)
    assert app.state.job_store.get(job_id).status == "WAITING_FOR_REVIEW"
    assert client.get(f"/api/polish/jobs/{job_id}/revisions").status_code == 200
    app.state.job_runner.shutdown()


def test_translation_review_uses_original_for_rejected_page(tmp_path: Path, monkeypatch):
    app, client = _application(tmp_path)
    pages = [ParsedPage(
        page_number=number, source_blocks=({"block_id": f"p{number}-b001", "kind": "text", "text": f"Source {number} mg."},),
        parser_backend="mineru", parser_version="test", ocr_used=False, figures=(), tables=(), equations=(), unresolved_items=(),
    ) for number in (1, 2)]
    monkeypatch.setattr("src.webapp.translation_service.parse_pdf_pages", lambda *args, **kwargs: pages)

    def build(**settings):
        def invoke(messages, **kwargs):
            source = json.loads(messages[1]["content"])["task_context"]["pages"][0]
            block = source["source_blocks"][0]
            return {
                "schema_version": "translation_page_v1",
                "pages": [{"page_number": source["page_number"], "translated_blocks": [
                    {"block_id": block["block_id"], "text": block["text"].replace("Source", "译文")},
                ], "unresolved_items": []}],
                "terminology_ledger": [], "unresolved_items": [],
            }
        return invoke

    monkeypatch.setattr("src.webapp.translation_service.build_web_json_llm_call", build)
    uploaded = client.post("/api/files", files={"file": ("paper.pdf", b"%PDF-1.7\nparser fixture", "application/pdf")})
    created = client.post("/api/translation/jobs", json={"input_file_id": uploaded.json()["file_id"]})
    job_id = created.json()["job_id"]
    app.state.job_runner._execute(job_id)
    assert app.state.job_store.get(job_id).status == "WAITING_FOR_REVIEW"
    rejected = client.post(f"/api/translation/jobs/{job_id}/actions", json={"type": "reject_pages", "page_numbers": [2]})
    assert rejected.status_code == 202, rejected.text
    assert client.post(f"/api/translation/jobs/{job_id}/actions", json={"type": "accept"}).status_code == 202
    exported = client.post(f"/api/translation/jobs/{job_id}/export", json={"format": "json"})
    assert exported.status_code == 200, exported.text
    document = exported.json()
    assert "译文 1 mg." in document["content"]
    assert "Source 2 mg." in document["content"]
    assert "译文 2 mg." not in document["content"]
    assert document["pages"][1]["review_status"] == "rejected"
    app.state.job_runner.shutdown()


def test_pdf_export_wraps_unicode_without_truncation():
    from pypdf import PdfReader

    content = "中文研究结果。" * 100 + "TAIL_SENTINEL"
    body, media_type, _filename = render_export(content, "pdf", title="Unicode document")
    assert media_type == "application/pdf"
    extracted = "".join(page.extract_text() for page in PdfReader(BytesIO(body)).pages)
    assert "中文研究结果" in extracted
    assert "TAIL_SENTINEL" in extracted


def test_translation_locks_restore_exact_anchors_and_reject_inventions():
    from src.webapp.page_blocks import protect_block, restore_block

    protected, anchors = protect_block("p1-b001", "A dose of 5 mg was used in Figure 2 [3].")
    translated = protected.replace("A dose of", "剂量为").replace("was used in", "用于")
    assert "5 mg" in restore_block(translated, anchors)
    with pytest.raises(ValueError, match="introduced"):
        restore_block(translated + " 99 mg", anchors)
    with pytest.raises(ValueError, match="changed"):
        restore_block(translated + next(iter(anchors)), anchors)


def test_mineru_content_list_preserves_page_and_block_mapping(tmp_path: Path, monkeypatch):
    from src.webapp.pdf_ingestion import _pages_from_mineru_content

    pages = _pages_from_mineru_content(
        [
            {"page_idx": 0, "type": "text", "text": "First page."},
            {"page_idx": 1, "type": "image", "image_caption": ["Figure 2"], "img_path": "fig.png"},
            {"page_idx": 1, "type": "table", "table_caption": ["Table 1"], "table_body": "a | b"},
        ],
        max_pages=2,
        parser_version="test",
        ocr_used=True,
    )
    assert [page.page_number for page in pages] == [1, 2]
    assert pages[0].source_blocks[0]["block_id"] == "p1-b001"
    assert pages[1].source_blocks[0]["kind"] == "image"
    assert "FIGURE" in str(pages[1].source_blocks[0]["text"])
    assert pages[1].source_blocks[1]["kind"] == "caption"
    assert pages[1].source_blocks[2]["kind"] == "table"
    assert pages[1].ocr_used is True


def test_selective_polish_regeneration_retains_accepted_segments(tmp_path: Path, monkeypatch):
    calls = _polish_provider(monkeypatch, lambda text: text.replace("observed", "measured"))
    app, client = _application(tmp_path)
    created = client.post("/api/polish/jobs", json={"text": "First was observed at 5 mg.\n\nSecond was observed at 8 mg."})
    job_id = created.json()["job_id"]
    app.state.job_runner._execute(job_id)
    assert client.post(f"/api/polish/jobs/{job_id}/actions", json={"type": "accept_segments", "segment_ids": ["p-0001"]}).status_code == 202
    regenerated = client.post(f"/api/polish/jobs/{job_id}/actions", json={
        "type": "regenerate_segments", "segment_ids": ["p-0002"], "notes": "Clarify only the second paragraph.",
        "expected_attempt": 1, "expected_review_version": 1,
    })
    assert regenerated.status_code == 202, regenerated.text
    assert regenerated.json()["revision"] == 2
    app.state.job_runner._execute(job_id)
    assert app.state.job_store.get(job_id).status == "WAITING_FOR_REVIEW"
    assert [call["context"]["segments"][0]["segment_id"] for call in calls] == ["p-0001", "p-0002", "p-0002"]
    current = client.get(f"/api/polish/jobs/{job_id}/revisions").json()
    assert current["review"]["segments"] == {"p-0001": "accepted"}
    assert current["review"]["version"] == 2
    assert calls[-1]["context"]["review_feedback"]["notes"] == "Clarify only the second paragraph."
    assert "previous_result" not in calls[-1]["context"]["review_feedback"]
    accepted = client.post(f"/api/polish/jobs/{job_id}/actions", json={"type": "accept", "expected_review_version": 2})
    assert accepted.status_code == 202
    exported = client.post(f"/api/polish/jobs/{job_id}/export", json={})
    assert exported.text == "First was measured at 5 mg.\n\nSecond was measured at 8 mg."
    app.state.job_runner.shutdown()


def test_selective_translation_regenerates_only_requested_page(tmp_path: Path, monkeypatch):
    app, client = _application(tmp_path)
    pages = [ParsedPage(
        number, ({"block_id": f"p{number}-b001", "kind": "text", "text": f"Source {number} mg."},),
        "mineru", "test", False, (), (), (), (),
    ) for number in (1, 2)]
    parse_calls = []
    calls = []

    def parse(*_args, **_kwargs):
        parse_calls.append(True)
        return pages

    def build(**_settings):
        def invoke(messages, **_kwargs):
            page = json.loads(messages[1]["content"])["task_context"]["pages"][0]
            calls.append(page["page_number"])
            block = page["source_blocks"][0]
            return {
                "schema_version": "translation_page_v1", "pages": [{
                    "page_number": page["page_number"], "translated_blocks": [{"block_id": block["block_id"], "text": block["text"].replace("Source", "译文")}],
                    "unresolved_items": [],
                }], "terminology_ledger": [], "unresolved_items": [],
            }
        return invoke

    monkeypatch.setattr("src.webapp.translation_service.parse_pdf_pages", parse)
    monkeypatch.setattr("src.webapp.translation_service.build_web_json_llm_call", build)
    uploaded = client.post("/api/files", files={"file": ("paper.pdf", b"%PDF-1.7\nparser fixture", "application/pdf")})
    created = client.post("/api/translation/jobs", json={"input_file_id": uploaded.json()["file_id"]})
    job_id = created.json()["job_id"]
    app.state.job_runner._execute(job_id)
    assert client.post(f"/api/translation/jobs/{job_id}/actions", json={"type": "accept_pages", "page_numbers": [1]}).status_code == 202
    response = client.post(f"/api/translation/jobs/{job_id}/actions", json={"type": "regenerate_pages", "page_numbers": [2], "notes": "Improve terminology."})
    assert response.status_code == 202, response.text
    app.state.job_runner._execute(job_id)
    assert app.state.job_store.get(job_id).status == "WAITING_FOR_REVIEW", app.state.job_store.get(job_id).error_message
    assert calls == [1, 2, 2]
    assert len(parse_calls) == 1
    summaries = client.get(f"/api/translation/jobs/{job_id}/pages").json()
    assert summaries["pages"][0]["review_status"] == "accepted"
    assert client.get(f"/api/translation/jobs/{job_id}/pages/1").json()["review_status"] == "accepted"
    app.state.job_runner.shutdown()


def test_empty_mineru_pages_fall_back_to_selectable_text(tmp_path: Path, monkeypatch):
    import src.webapp.pdf_ingestion as ingestion
    from reportlab.pdfgen.canvas import Canvas

    source = tmp_path / "source.pdf"
    canvas = Canvas(str(source))
    canvas.drawString(50, 750, "Selectable evidence remains available.")
    canvas.save()
    original_find = ingestion.importlib.util.find_spec
    monkeypatch.setattr(ingestion.importlib.util, "find_spec", lambda name: object() if name == "mineru" else original_find(name))

    def empty_mineru(*_args, **kwargs):
        kwargs["diagnostics"].append("mineru_empty_or_invalid_content_list")
        return [ParsedPage(1, (), "mineru", "test", False, (), (), (), ())]

    monkeypatch.setattr(ingestion, "_try_mineru", empty_mineru)
    pages = ingestion.parse_pdf_pages(source, work_dir=tmp_path / "parsed")
    assert pages[0].parser_backend == "pypdf"
    assert "Selectable evidence" in pages[0].source_blocks[0]["text"]
    assert "mineru_empty_or_invalid_content_list" in pages[0].parser_diagnostics


def test_page_id_indices_are_normalized_without_guessing_ambiguous_pages():
    from src.webapp.pdf_ingestion import _pages_from_mineru_content

    settings = {"max_pages": 2, "parser_version": "test", "ocr_used": False}
    pages = _pages_from_mineru_content([{"page_id": 0, "text": "First"}, {"page_id": 1, "text": "Second"}], **settings)
    assert [page.source_blocks[0]["text"] for page in pages] == ["First", "Second"]
    assert _pages_from_mineru_content([{"page_id": 1, "text": "Unknown base"}], **settings) == []


def test_polish_preserves_docx_body_order_and_structured_content(tmp_path: Path, monkeypatch):
    from docx import Document
    from src.webapp.worker_support import source_text

    _polish_provider(monkeypatch, lambda text: text.replace("observed", "measured"))
    app, client = _application(tmp_path)
    document = Document()
    document.add_paragraph("Before was observed.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Table was observed at 5 mg."
    document.add_paragraph("After was observed.")
    buffer = BytesIO()
    document.save(buffer)
    uploaded = client.post("/api/files", files={"file": ("draft.docx", buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    file_id = uploaded.json()["file_id"]
    record = app.state.file_service.get(file_id)
    text = source_text(app.state.file_service, {"files": [record.model_dump()]})
    assert text.index("Before") < text.index("Table") < text.index("After")
    created = client.post("/api/polish/jobs", json={"input_file_id": file_id})
    job_id = created.json()["job_id"]
    app.state.job_runner._execute(job_id)
    result = client.get(f"/api/polish/jobs/{job_id}/revisions").json()
    assert result["segments"][1]["revised"] == "| Table was observed at 5 mg. |"
    app.state.job_runner.shutdown()
