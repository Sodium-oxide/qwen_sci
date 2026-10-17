"""Reviewed document assembly and content-addressed final exports."""

from __future__ import annotations

import hashlib
import json
import re
from copy import deepcopy
from io import BytesIO
from typing import Any, Mapping
from xml.sax.saxutils import escape

from filelock import FileLock

from .document_segments import segment_text
from .feature_review import read_feature_result, read_review
from .file_service import FileService
from .job_artifacts import publish_export_artifact
from .jobs import JobConflictError, JobStateError, JobStore
from .jobs.models import JobRecord
from .worker_support import source_text


class ExportError(ValueError):
    code = "EXPORT_INVALID"


class ExportUnavailable(ExportError):
    code = "CAPABILITY_UNAVAILABLE"


def _manifest(store: JobStore, record: JobRecord) -> dict[str, Any]:
    try:
        raw = (store.job_directory(record.job_id) / "input.json").read_bytes()
        payload = json.loads(raw)
    except (OSError, UnicodeError, ValueError) as exc:
        raise ExportError("The frozen export input is unavailable.") from exc
    if not isinstance(payload, dict) or hashlib.sha256(raw).hexdigest() != record.input_manifest_sha256:
        raise ExportError("The frozen export input no longer matches its recorded hash.")
    return payload


def export_polish_document(
    store: JobStore, job: JobRecord, manifest: Mapping[str, Any], file_service: FileService,
) -> tuple[str, dict[str, Any]]:
    candidate = read_feature_result(store, job, "polish_revision_set")
    revisions = candidate.get("segments")
    if not isinstance(revisions, list):
        raise ExportError("The polishing result has no revision segments.")
    original = source_text(file_service, manifest["input"])
    segments = segment_text(original)
    by_id = {item.get("segment_id"): item for item in revisions if isinstance(item, dict)}
    if len(by_id) != len(revisions) or set(by_id) != {segment.segment_id for segment in segments}:
        raise ExportError("The revision segments do not match the frozen source.")
    decisions = read_review(store, job).get("segments", {})
    cursor = 0
    chunks = []
    reviewed = []
    for segment in segments:
        revision = by_id[segment.segment_id]
        if revision.get("original") != segment.text or not isinstance(revision.get("revised"), str):
            raise ExportError("A revision no longer matches its source segment.")
        accepted = decisions.get(segment.segment_id) != "rejected"
        selected = revision["revised"] if accepted else segment.text
        chunks.extend((original[cursor:segment.start], selected))
        cursor = segment.end
        reviewed.append({**revision, "review_status": "accepted" if accepted else "rejected", "selected_text": selected})
    chunks.append(original[cursor:])
    content = "".join(chunks)
    return content, {**candidate, "segments": reviewed, "content": content}


def export_translation_document(store: JobStore, job: JobRecord) -> tuple[str, dict[str, Any]]:
    candidate = deepcopy(read_feature_result(store, job, "translation_pages"))
    pages = candidate.get("pages")
    if not isinstance(pages, list):
        raise ExportError("The translation result has no pages.")
    decisions = read_review(store, job).get("pages", {})
    chunks = []
    for page in pages:
        page_number = page["page_number"]
        rejected = decisions.get(str(page_number)) == "rejected"
        page["review_status"] = "rejected" if rejected else "accepted"
        page["selected_blocks"] = page["source_blocks"] if rejected else page["translated_blocks"]
        chunks.append(f"## Page {page_number}")
        if rejected:
            chunks.append("[Translation rejected; original source retained.]")
        chunks.extend(str(block["text"]) for block in page["selected_blocks"])
        if page.get("unresolved_items"):
            chunks.append("[AUTHOR_INPUT_NEEDED: review unresolved source or terminology items.]")
    content = "\n\n".join(chunks).strip() + "\n"
    candidate["content"] = content
    return content, candidate


def export_feature(
    store: JobStore, file_service: FileService, *, kind: str, job_id: str,
    output_format: str, expected_review_version: int | None = None,
) -> tuple[bytes, str, str, str]:
    directory = store.job_directory(job_id)
    with FileLock(str(directory / ".review.guard"), timeout=10):
        record = store.get(job_id)
        if record.kind != kind or kind not in {"polish", "translation"}:
            raise ExportError("This document export does not apply to the selected Job.")
        if record.status != "SUCCEEDED":
            raise JobStateError("Accept the Job review before exporting the final document.")
        review = read_review(store, record)
        if review.get("decision") != "accept":
            raise JobStateError("This Job has no accepted review for its current attempt.")
        if expected_review_version is not None and expected_review_version != review["version"]:
            raise JobConflictError("The requested export review version has changed.")
        manifest = _manifest(store, record)
        if kind == "polish":
            content, payload = export_polish_document(store, record, manifest, file_service)
        else:
            content, payload = export_translation_document(store, record)
        payload["review"] = review
        payload["job_id"] = job_id
        payload["attempt"] = record.attempt
        body, media_type, filename = render_export(content, output_format, title=f"{kind}-{job_id}", payload=payload)
        artifact_id = publish_export_artifact(
            store, record, filename=filename, content=body, media_type=media_type,
            provenance={
                "review_version": review["version"], "exporter_version": "2.0.0",
                "input_manifest_sha256": record.input_manifest_sha256,
                "prompt_hash": record.prompt_hash,
            },
            source_file_ids=[item["file_id"] for item in manifest["input"].get("files", [])],
        )
        return body, media_type, filename, artifact_id


def render_export(
    content: str, output_format: str, *, title: str, payload: Mapping[str, Any] | None = None,
) -> tuple[bytes, str, str]:
    if output_format == "markdown":
        return content.encode("utf-8"), "text/markdown; charset=utf-8", f"{title}.md"
    if output_format == "json":
        return (
            json.dumps(payload or {"content": content}, ensure_ascii=False, indent=2, allow_nan=False).encode("utf-8") + b"\n",
            "application/json", f"{title}.json",
        )
    if output_format == "docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ExportUnavailable("DOCX export requires python-docx.") from exc
        document = Document()
        for block in re.split(r"\n\s*\n", content):
            heading = re.fullmatch(r"(#{1,6})\s+(.+)", block)
            if heading:
                document.add_heading(heading.group(2), level=len(heading.group(1)))
            else:
                document.add_paragraph(block)
        buffer = BytesIO()
        document.save(buffer)
        return buffer.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document", f"{title}.docx"
    if output_format == "pdf":
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import ParagraphStyle
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        except ImportError as exc:
            raise ExportUnavailable("PDF export requires reportlab.") from exc
        if "STSong-Light" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        style = ParagraphStyle("Document", fontName="STSong-Light", fontSize=11, leading=17, wordWrap="CJK", splitLongWords=True)
        buffer = BytesIO()
        document = SimpleDocTemplate(buffer, pagesize=A4, title=title, leftMargin=48, rightMargin=48, topMargin=48, bottomMargin=48)
        story = []
        for block in re.split(r"\n\s*\n", content):
            story.extend((Paragraph(escape(block).replace("\n", "<br/>"), style), Spacer(1, 8)))
        document.build(story)
        return buffer.getvalue(), "application/pdf", f"{title}.pdf"
    raise ExportError("Unsupported export format.")
